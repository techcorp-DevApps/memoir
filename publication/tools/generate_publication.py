#!/usr/bin/env python3
from __future__ import annotations
import csv, json, re, shutil, subprocess, tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import List, Tuple
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

ROOT = Path(__file__).resolve().parents[2]
CHAPTER_DIR = ROOT / 'chapters'
OUT_DIR = ROOT / 'publication' / 'chapters'
QA_DIR = ROOT / 'publication' / 'qa'
MANIFEST_CSV = ROOT / 'publication' / 'publication-formatting-manifest.csv'
MANIFEST_MD = ROOT / 'publication' / 'publication-formatting-manifest.md'

BLOCKED = {
    '12_TheOnion.md': 'BLOCKED — exact author-approved consolidated Onion/Stockpot source required by 2026-09-15 publication instruction is not present in repository',
    '13_TheStockPot.md': 'BLOCKED — superseded as standalone chapter by 2026-09-15 publication instruction; must not be formatted independently',
}
EXCLUDED = {
    '23_Dave.md': 'EXCLUDED — current planning record classifies file as placeholder, not complete manuscript chapter',
    '24_ReturnToYulara.md': 'EXCLUDED — current planning record classifies file as placeholder; developed draft remains elsewhere and is not promoted',
}

@dataclass
class Block:
    kind: str
    text: str
    level: int = 0

INLINE_RE = re.compile(r'(\*\*.+?\*\*|__.+?__|\*[^*\n]+?\*|_[^_\n]+?_|`[^`\n]+?`|\[[^\]]+\]\([^\)]+\))')

def run(cmd: List[str], cwd: Path | None = None) -> subprocess.CompletedProcess:
    return subprocess.run(cmd, cwd=str(cwd) if cwd else None, text=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE, check=True)

def git_sha(path: Path) -> str:
    return run(['git','rev-parse','HEAD'], ROOT).stdout.strip()

def git_blob_sha(path: Path) -> str:
    rel = str(path.relative_to(ROOT))
    return run(['git','rev-parse',f'HEAD:{rel}'], ROOT).stdout.strip()

def parse_markdown(text: str) -> List[Block]:
    lines = text.replace('\r\n','\n').replace('\r','\n').split('\n')
    blocks: List[Block] = []
    i = 0
    while i < len(lines):
        s = lines[i].strip()
        if not s:
            i += 1; continue
        if s == '---':
            # Repository convention: a rule immediately after the opening H1 is a
            # Markdown title divider, not an in-scene manuscript break. Publication
            # layout already supplies title spacing, so suppress only that initial
            # divider; all later rules are rendered as * * *.
            if blocks and len(blocks) == 1 and blocks[0].kind == 'heading' and blocks[0].level == 1:
                i += 1; continue
            blocks.append(Block('scene',''))
            i += 1; continue
        m = re.match(r'^(#{1,6})\s+(.*)$', s)
        if m:
            blocks.append(Block('heading',m.group(2).strip(),len(m.group(1))))
            i += 1; continue
        if s.startswith('>'):
            parts=[]
            while i < len(lines) and lines[i].lstrip().startswith('>'):
                v = lines[i].lstrip()[1:]
                if v.startswith(' '): v=v[1:]
                parts.append(v.strip())
                i += 1
            blocks.append(Block('quote',' '.join(parts).strip()))
            continue
        m = re.match(r'^[-*+]\s+(.*)$', s)
        if m:
            item = m.group(1).strip(); i += 1
            while i < len(lines) and lines[i].strip() and not re.match(r'^(?:[-*+]\s+|\d+[.)]\s+|#{1,6}\s+|>|---$)', lines[i].strip()):
                item += ' ' + lines[i].strip(); i += 1
            blocks.append(Block('bullet',item)); continue
        m = re.match(r'^\d+[.)]\s+(.*)$', s)
        if m:
            item = m.group(1).strip(); i += 1
            while i < len(lines) and lines[i].strip() and not re.match(r'^(?:[-*+]\s+|\d+[.)]\s+|#{1,6}\s+|>|---$)', lines[i].strip()):
                item += ' ' + lines[i].strip(); i += 1
            blocks.append(Block('number',item)); continue
        parts=[s]; i += 1
        while i < len(lines):
            t=lines[i].strip()
            if not t: break
            if t == '---' or re.match(r'^(?:#{1,6}\s+|>|[-*+]\s+|\d+[.)]\s+)', t): break
            parts.append(t); i += 1
        blocks.append(Block('body',' '.join(parts)))
    return blocks

def visible_text(md: str) -> str:
    def repl(m):
        tok=m.group(0)
        if tok.startswith('**') or tok.startswith('__'): return tok[2:-2]
        if tok.startswith('*') or tok.startswith('_') or tok.startswith('`'): return tok[1:-1]
        if tok.startswith('['): return tok[1:tok.index(']')]
        return tok
    return INLINE_RE.sub(repl, md)

def add_inline_runs(p, text: str):
    pos=0
    for m in INLINE_RE.finditer(text):
        if m.start() > pos: p.add_run(text[pos:m.start()])
        tok=m.group(0)
        if tok.startswith('**') or tok.startswith('__'):
            r=p.add_run(tok[2:-2]); r.bold=True
        elif tok.startswith('*') or tok.startswith('_'):
            r=p.add_run(tok[1:-1]); r.italic=True
        elif tok.startswith('`'):
            p.add_run(tok[1:-1])
        elif tok.startswith('['):
            p.add_run(tok[1:tok.index(']')])
        pos=m.end()
    if pos < len(text): p.add_run(text[pos:])

def set_font(run, name='EB Garamond', size=11.5):
    run.font.name=name; run.font.size=Pt(size)
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)

def style_runs(p, size=11.5):
    for r in p.runs: set_font(r, size=size)

def widow_control(p):
    pPr=p._p.get_or_add_pPr(); pPr.append(OxmlElement('w:widowControl'))

def keep_with_next(p):
    pPr=p._p.get_or_add_pPr(); pPr.append(OxmlElement('w:keepNext'))

def set_mirror_margins(doc: Document):
    settings=doc.settings._element
    if settings.find(qn('w:mirrorMargins')) is None:
        settings.append(OxmlElement('w:mirrorMargins'))

def set_doc_defaults(doc: Document):
    sec=doc.sections[0]
    sec.page_width=Inches(6); sec.page_height=Inches(9)
    sec.top_margin=Inches(.72); sec.bottom_margin=Inches(.68)
    sec.left_margin=Inches(.82); sec.right_margin=Inches(.62)
    sec.header_distance=Inches(.30); sec.footer_distance=Inches(.32)
    set_mirror_margins(doc)
    normal=doc.styles['Normal']
    normal.font.name='EB Garamond'; normal.font.size=Pt(11.5)
    normal._element.rPr.rFonts.set(qn('w:eastAsia'),'EB Garamond')
    pf=normal.paragraph_format
    pf.space_before=Pt(0); pf.space_after=Pt(0); pf.line_spacing=Pt(14)

def format_doc(blocks: List[Block], out_path: Path) -> Tuple[str,str]:
    doc=Document(); set_doc_defaults(doc)
    first_body_after_break=True
    seen_title=False
    prev_was_cheftip=False
    for b in blocks:
        if b.kind=='scene':
            p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before=Pt(14); p.paragraph_format.space_after=Pt(14)
            r=p.add_run('* * *'); set_font(r,size=10.5)
            keep_with_next(p); first_body_after_break=True; prev_was_cheftip=False
            continue
        if b.kind=='heading':
            p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            if b.level==1 and not seen_title:
                p.paragraph_format.space_before=Pt(66); p.paragraph_format.space_after=Pt(34)
                add_inline_runs(p,b.text); style_runs(p,15.5)
                for r in p.runs: r.bold=True
                seen_title=True; first_body_after_break=True; prev_was_cheftip=False
            elif b.level==2 and visible_text(b.text).strip().lower().startswith('cheftip #'):
                p.paragraph_format.space_before=Pt(18); p.paragraph_format.space_after=Pt(5)
                add_inline_runs(p,b.text); style_runs(p,12.5)
                for r in p.runs: r.bold=True
                keep_with_next(p); first_body_after_break=True; prev_was_cheftip=True
            else:
                p.paragraph_format.space_before=Pt(18); p.paragraph_format.space_after=Pt(10)
                add_inline_runs(p,b.text); style_runs(p,12.5)
                for r in p.runs: r.bold=True
                keep_with_next(p); first_body_after_break=True; prev_was_cheftip=False
            continue
        p=doc.add_paragraph(); widow_control(p)
        if prev_was_cheftip and b.kind in ('body','quote'):
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.left_indent=Inches(.22); p.paragraph_format.right_indent=Inches(.22)
            p.paragraph_format.first_line_indent=Inches(0); p.paragraph_format.space_after=Pt(9)
            add_inline_runs(p,b.text); style_runs(p,11.5)
            for r in p.runs: r.italic=True
            prev_was_cheftip=False; first_body_after_break=False
            continue
        if b.kind=='quote':
            p.paragraph_format.left_indent=Inches(.22); p.paragraph_format.right_indent=Inches(.22)
            p.paragraph_format.first_line_indent=Inches(0)
            p.paragraph_format.space_before=Pt(4); p.paragraph_format.space_after=Pt(4)
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        elif b.kind in ('bullet','number'):
            p.paragraph_format.left_indent=Inches(.22); p.paragraph_format.first_line_indent=Inches(-.14)
            p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
            p.add_run('• ')
        else:
            p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.first_line_indent=Inches(0 if first_body_after_break else .22)
        p.paragraph_format.line_spacing=Pt(14); p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(0)
        add_inline_runs(p,b.text); style_runs(p,11.5)
        prev_was_cheftip=False; first_body_after_break=False
    if doc.paragraphs and doc.paragraphs[0].text=='' and len(doc.paragraphs)>1:
        el=doc.paragraphs[0]._element; el.getparent().remove(el)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    src='\n'.join(visible_text(b.text) for b in blocks if b.kind!='scene')
    dst='\n'.join(re.sub(r'^•\s+','',p.text) for p in doc.paragraphs if p.text.strip()!='* * *')
    return src,dst

def normalize(s: str) -> str:
    s=s.replace('•',' ').replace('\u00ad','')
    return re.sub(r'\s+',' ',s).strip()

def convert_pdf(docx: Path, outdir: Path) -> Path:
    outdir.mkdir(parents=True, exist_ok=True)
    run(['libreoffice','--headless','--convert-to','pdf','--outdir',str(outdir),str(docx)])
    pdf=outdir/(docx.stem+'.pdf')
    if not pdf.exists(): raise RuntimeError(f'PDF not created for {docx}')
    return pdf

def pdf_info(pdf: Path) -> Tuple[int,str]:
    info=run(['pdfinfo',str(pdf)]).stdout
    m=re.search(r'^Pages:\s+(\d+)',info,re.M); pages=int(m.group(1)) if m else 0
    fonts=run(['pdffonts',str(pdf)]).stdout
    return pages,fonts

def pdf_text(pdf: Path) -> str:
    with tempfile.NamedTemporaryFile(suffix='.txt',delete=False) as f: tmp=Path(f.name)
    try:
        run(['pdftotext',str(pdf),str(tmp)])
        return tmp.read_text('utf-8',errors='replace')
    finally:
        tmp.unlink(missing_ok=True)

def raster_qa(pdf: Path, chapter_key: str) -> Tuple[str,List[str],List[Path]]:
    d=QA_DIR/chapter_key
    if d.exists(): shutil.rmtree(d)
    d.mkdir(parents=True)
    prefix=d/'page'
    run(['pdftoppm','-png','-r','144',str(pdf),str(prefix)])
    pages=sorted(d.glob('page-*.png'))
    issues=[]
    expected=(864,1296)
    for idx,p in enumerate(pages,1):
        im=Image.open(p).convert('L')
        if im.size != expected: issues.append(f'p{idx}:unexpected_size={im.size}')
        mask=im.point(lambda x: 255 if x < 245 else 0)
        bbox=mask.getbbox()
        if bbox is None:
            issues.append(f'p{idx}:blank_page'); continue
        x0,y0,x1,y1=bbox
        if x0 < 12 or y0 < 12 or x1 > im.width-12 or y1 > im.height-12:
            issues.append(f'p{idx}:content_near_trim_edge bbox={bbox}')
    return ('PASS' if not issues else 'REVIEW'),issues,pages

def safe_name(num: str, title: str) -> str:
    t=re.sub(r'[^A-Za-z0-9]+','_',title.strip()).strip('_')
    return f'{num}_{t}'

def main():
    OUT_DIR.mkdir(parents=True,exist_ok=True); QA_DIR.mkdir(parents=True,exist_ok=True)
    rows=[]; total_pages=0
    for src in sorted(CHAPTER_DIR.glob('*.md')):
        name=src.name; num=name.split('_',1)[0]
        if name in EXCLUDED:
            rows.append({'order':num,'title':name,'source_path':str(src.relative_to(ROOT)),'source_blob_sha':git_blob_sha(src),'approval_status':'PLACEHOLDER','docx_output':'','pdf_output':'','page_count':'','docx_text_compare':'NOT_RUN','pdf_text_compare':'NOT_RUN','render_preflight':'NOT_RUN','visual_qa':'NOT_RUN','concern':EXCLUDED[name],'sync_status':'branch-generated'}); continue
        if name in BLOCKED:
            rows.append({'order':num,'title':name,'source_path':str(src.relative_to(ROOT)),'source_blob_sha':git_blob_sha(src),'approval_status':'BLOCKED_SOURCE_CONFLICT','docx_output':'','pdf_output':'','page_count':'','docx_text_compare':'NOT_RUN','pdf_text_compare':'NOT_RUN','render_preflight':'NOT_RUN','visual_qa':'NOT_RUN','concern':BLOCKED[name],'sync_status':'blocked'}); continue
        text=src.read_text('utf-8'); blocks=parse_markdown(text)
        title=next((visible_text(b.text) for b in blocks if b.kind=='heading' and b.level==1),name)
        base=safe_name(num,title); docx=OUT_DIR/(base+'.docx')
        src_txt,docx_txt=format_doc(blocks,docx)
        docx_cmp='PASS' if normalize(src_txt)==normalize(docx_txt) else 'FAIL'
        pdf=convert_pdf(docx,OUT_DIR); pages,fonts=pdf_info(pdf); total_pages += pages
        ptxt=pdf_text(pdf)
        pdf_cmp='PASS' if normalize(src_txt) in normalize(ptxt) or normalize(src_txt)==normalize(ptxt) else 'REVIEW'
        render_status,issues,_=raster_qa(pdf,base)
        font_status='PASS' if re.search(r'EB.?Garamond',fonts,re.I) else 'REVIEW'
        concern='; '.join(issues)
        if name=='03_FreshForUnlock.md':
            concern=(concern+'; ' if concern else '')+'source contains additional H1 GENERAL POPULATION; rendered as internal centred heading; structural status remains for author review'
        rows.append({'order':num,'title':title,'source_path':str(src.relative_to(ROOT)),'source_blob_sha':git_blob_sha(src),'approval_status':'CURRENT_REPOSITORY_SOURCE','docx_output':str(docx.relative_to(ROOT)),'pdf_output':str(pdf.relative_to(ROOT)),'page_count':pages,'docx_text_compare':docx_cmp,'pdf_text_compare':pdf_cmp,'render_preflight':f'{render_status}; font={font_status}','visual_qa':'PENDING_HUMAN_100_PERCENT','concern':concern,'sync_status':'branch-generated'})
    fields=['order','title','source_path','source_blob_sha','approval_status','docx_output','pdf_output','page_count','docx_text_compare','pdf_text_compare','render_preflight','visual_qa','concern','sync_status']
    MANIFEST_CSV.parent.mkdir(parents=True,exist_ok=True)
    with MANIFEST_CSV.open('w',encoding='utf-8',newline='') as f:
        w=csv.DictWriter(f,fieldnames=fields); w.writeheader(); w.writerows(rows)
    commit=git_sha(ROOT)
    lines=['# Publication Formatting Manifest','',f'- Source commit: `{commit}`',f'- Generated chapters: {sum(1 for r in rows if r["docx_output"])}',f'- Generated pages: {total_pages}','- Individual chapter folios: none','- Human 100% visual QA: **pending**; automated raster preflight is recorded separately.','- Source conflict: consolidated THE ONION WILL NOT WAIT FOR YOU / Stockpot text required by the 2026-09-15 publication instruction is not present as an exact authoritative source in the repository, so the two older standalone files are intentionally not emitted.','', '| Order | Title | Source | Pages | DOCX text | PDF text | Raster preflight | Visual QA | Concern |','|---:|---|---|---:|---|---|---|---|---|']
    for r in rows:
        title=str(r['title']).replace('|',' / '); concern=str(r['concern']).replace('|',' / ')
        lines.append(f'| {r["order"]} | {title} | `{r["source_path"]}` | {r["page_count"]} | {r["docx_text_compare"]} | {r["pdf_text_compare"]} | {r["render_preflight"]} | {r["visual_qa"]} | {concern} |')
    MANIFEST_MD.write_text('\n'.join(lines)+'\n','utf-8')
    print(json.dumps({'commit':commit,'generated':sum(1 for r in rows if r['docx_output']),'total_pages':total_pages,'manifest':str(MANIFEST_MD.relative_to(ROOT))},indent=2))

if __name__=='__main__': main()
